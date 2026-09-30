import sys
import time
import queue
import threading
from datetime import datetime

import requests
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

# ============================== KONFIGURASI ==============================
FILE_EXCEL = "Jawa Tengah.xlsx"
JUMLAH_THREAD = 8
TIMEOUT = 20            # detik per permintaan HTTP
MAKS_PERCOBAAN = 4      # jumlah retry bila gagal / kena rate limit
USER_AGENT = "TugasMultithreadingCuaca/1.0 (tugas kuliah)"  # wajib untuk Nominatim

URL_GEOCODING = "https://geocoding-api.open-meteo.com/v1/search"
URL_NOMINATIM = "https://nominatim.openstreetmap.org/search"
URL_CUACA = "https://api.open-meteo.com/v1/forecast"

NAMA_PROVINSI = ("central java", "jawa tengah")
# Kotak batas kasar Jawa Tengah: membuang hasil geocoding yang nyasar ke daerah lain
LAT_MIN, LAT_MAX, LON_MIN, LON_MAX = -8.9, -5.6, 108.4, 111.9

# Nama kolom harus sama persis dengan header di Excel
KOLOM_CUACA = ["Last Update", "Suhu (°C)", "Kelembapan (%)", "Kondisi Cuaca",
               "Kecepatan Angin (km/jam)", "Arah Angin", "Sinar UV"]
KOLOM_TAMBAHAN = ["Latitude", "Longitude", "Catatan"]

# Kode cuaca WMO -> Bahasa Indonesia
KODE_CUACA = {
    0: "Cerah", 1: "Cerah Berawan", 2: "Berawan Sebagian", 3: "Berawan",
    45: "Berkabut", 48: "Kabut Beku",
    51: "Gerimis Ringan", 53: "Gerimis Sedang", 55: "Gerimis Lebat",
    56: "Gerimis Beku Ringan", 57: "Gerimis Beku Lebat",
    61: "Hujan Ringan", 63: "Hujan Sedang", 65: "Hujan Lebat",
    66: "Hujan Beku Ringan", 67: "Hujan Beku Lebat",
    71: "Salju Ringan", 73: "Salju Sedang", 75: "Salju Lebat", 77: "Butiran Salju",
    80: "Hujan Lokal Ringan", 81: "Hujan Lokal Sedang", 82: "Hujan Lokal Lebat",
    85: "Hujan Salju Ringan", 86: "Hujan Salju Lebat",
    95: "Badai Petir", 96: "Badai Petir + Hujan Es Ringan", 99: "Badai Petir + Hujan Es Lebat",
}
ARAH_MATA_ANGIN = ["Utara", "Timur Laut", "Timur", "Tenggara",
                   "Selatan", "Barat Daya", "Barat", "Barat Laut"]

# Nominatim membatasi ~1 permintaan/detik, jadi semua thread berbagi satu lock
_lock_nominatim = threading.Lock()
_terakhir_nominatim = [0.0]


# ============================== FUNGSI BANTU ==============================
def derajat_ke_arah(derajat):
    """135 -> 'Tenggara (135°)' (8 arah mata angin, tiap arah selebar 45°)."""
    if derajat is None:
        return "-"
    return f"{ARAH_MATA_ANGIN[int((derajat + 22.5) // 45) % 8]} ({round(derajat)}°)"


def format_waktu(iso_string):
    """'2026-09-30T14:15' -> '30-09-2026 14:15 WIB'."""
    try:
        return datetime.fromisoformat(iso_string).strftime("%d-%m-%Y %H:%M") + " WIB"
    except (TypeError, ValueError):
        return iso_string or "-"


def bersihkan_nama(teks):
    """Buang awalan 'Kabupaten'/'Kota'/'Kecamatan' agar cocok dengan hasil geocoding."""
    t = (teks or "").strip()
    for awalan in ("kabupaten ", "kab. ", "kota ", "kecamatan ", "kec. "):
        if t.lower().startswith(awalan):
            t = t[len(awalan):]
    return t.strip()


def dalam_jateng(lat, lon):
    return LAT_MIN <= lat <= LAT_MAX and LON_MIN <= lon <= LON_MAX


def get_json(url, params, headers=None):
    """GET + retry dengan exponential backoff (1,5 s, 3 s, 6 s, ...) untuk 429/5xx."""
    kesalahan = None
    for percobaan in range(MAKS_PERCOBAAN):
        try:
            r = requests.get(url, params=params, headers=headers, timeout=TIMEOUT)
            if r.status_code == 200:
                return r.json()
            if r.status_code in (429, 500, 502, 503, 504):
                kesalahan = f"HTTP {r.status_code}"
            else:
                r.raise_for_status()
        except requests.RequestException as e:
            kesalahan = str(e)
        time.sleep(1.5 * (2 ** percobaan))
    raise RuntimeError(f"Gagal mengambil data: {kesalahan}")


# ============================== GEOCODING ==============================
def geocode_open_meteo(nama, kabupaten):
    """Cari koordinat di Open-Meteo. Kabupaten dicocokkan karena banyak nama kecamatan kembar."""
    data = get_json(URL_GEOCODING, {"name": nama, "count": 20, "language": "id",
                                    "format": "json", "countryCode": "ID"})
    kandidat = [h for h in (data.get("results") or [])
                if (h.get("admin1") or "").lower() in NAMA_PROVINSI
                and dalam_jateng(h["latitude"], h["longitude"])]
    if not kandidat:
        return None
    if kabupaten:
        for h in kandidat:
            if kabupaten.lower() in (h.get("admin2") or "").lower():
                return h["latitude"], h["longitude"]
        return None  # nama ada, tapi di kabupaten lain
    return kandidat[0]["latitude"], kandidat[0]["longitude"]


def geocode_nominatim(kecamatan, kabupaten_lengkap):
    """Cadangan: OpenStreetMap, dibatasi kotak Jawa Tengah."""
    with _lock_nominatim:  # pastikan jeda >= 1,1 detik antar permintaan
        jeda = 1.1 - (time.time() - _terakhir_nominatim[0])
        if jeda > 0:
            time.sleep(jeda)
        try:
            data = get_json(URL_NOMINATIM, {
                "q": f"Kecamatan {kecamatan}, {kabupaten_lengkap}, Jawa Tengah, Indonesia",
                "format": "json", "limit": 1, "countrycodes": "id",
                "viewbox": f"{LON_MIN},{LAT_MAX},{LON_MAX},{LAT_MIN}", "bounded": 1,
            }, headers={"User-Agent": USER_AGENT})
        finally:
            _terakhir_nominatim[0] = time.time()
    if data:
        lat, lon = float(data[0]["lat"]), float(data[0]["lon"])
        if dalam_jateng(lat, lon):
            return lat, lon
    return None


def cari_koordinat(kecamatan, kabupaten_lengkap):
    """Urutan: Open-Meteo -> Nominatim -> pusat kab/kota (pendekatan, dicatat di Excel)."""
    kec, kab = bersihkan_nama(kecamatan), bersihkan_nama(kabupaten_lengkap)

    hasil = geocode_open_meteo(kec, kab)
    if hasil:
        return hasil[0], hasil[1], ""
    hasil = geocode_nominatim(kec, kabupaten_lengkap)
    if hasil:
        return hasil[0], hasil[1], "Koordinat dari OpenStreetMap"
    hasil = geocode_open_meteo(kab, None)
    if hasil:
        return hasil[0], hasil[1], "Koordinat pendekatan (pusat kab/kota)"
    raise LookupError("Koordinat tidak ditemukan")


# ============================== DATA CUACA ==============================
def ambil_cuaca(lat, lon):
    """Ambil 7 data cuaca saat ini dari Open-Meteo Forecast API."""
    data = get_json(URL_CUACA, {
        "latitude": lat, "longitude": lon,
        "current": "temperature_2m,relative_humidity_2m,weather_code,"
                   "wind_speed_10m,wind_direction_10m,uv_index",
        "wind_speed_unit": "kmh", "timezone": "Asia/Jakarta",
    })
    cur = data["current"]

    uv = cur.get("uv_index")
    if uv is None:  # cadangan: ambil dari data per jam pada jam yang sama
        per_jam = get_json(URL_CUACA, {"latitude": lat, "longitude": lon,
                                       "hourly": "uv_index", "forecast_days": 1,
                                       "timezone": "Asia/Jakarta"})
        jam = cur["time"][:13]
        for t, v in zip(per_jam["hourly"]["time"], per_jam["hourly"]["uv_index"]):
            if t.startswith(jam):
                uv = v
                break

    return {
        "Last Update": format_waktu(cur.get("time")),
        "Suhu (°C)": cur.get("temperature_2m"),
        "Kelembapan (%)": cur.get("relative_humidity_2m"),
        "Kondisi Cuaca": KODE_CUACA.get(cur.get("weather_code"), f"Kode {cur.get('weather_code')}"),
        "Kecepatan Angin (km/jam)": cur.get("wind_speed_10m"),
        "Arah Angin": derajat_ke_arah(cur.get("wind_direction_10m")),
        "Sinar UV": uv,
    }


# ============================== MULTITHREADING ==============================
class PekerjaCuaca(threading.Thread):
    """
    Satu thread pekerja (consumer). Semua thread mengambil tugas (1 kecamatan)
    dari Queue yang sama sampai kosong. Hasil masuk ke dictionary bersama
    yang dilindungi Lock agar tidak terjadi race condition.
    """

    def __init__(self, nomor, antrean, hasil, lock, progres):
        super().__init__(name=f"Pekerja-{nomor}", daemon=True)
        self.antrean, self.hasil, self.lock, self.progres = antrean, hasil, lock, progres

    def run(self):
        while True:
            try:
                baris, kabupaten, kecamatan, lat, lon = self.antrean.get_nowait()
            except queue.Empty:
                return  # tugas habis -> thread selesai

            catatan = ""
            try:
                if lat is None or lon is None:  # koordinat manual di Excel dipakai bila ada
                    lat, lon, catatan = cari_koordinat(kecamatan, kabupaten)
                data = ambil_cuaca(lat, lon)
                data.update({"Latitude": lat, "Longitude": lon, "Catatan": catatan})
            except Exception as e:  # satu kecamatan gagal tidak menghentikan yang lain
                data = {"error": str(e), "Latitude": lat, "Longitude": lon}

            with self.lock:
                self.hasil[baris] = data
                self.progres["selesai"] += 1
                status = "GAGAL: " + data["error"] if "error" in data else "OK"
                print(f"[{self.progres['selesai']}/{self.progres['total']}] "
                      f"{self.name}: {kecamatan} ({kabupaten}) -> {status}")
            self.antrean.task_done()


# ============================== EXCEL ==============================
def peta_kolom(ws):
    """Petakan header -> nomor kolom (tidak bergantung urutan); buat kolom tambahan bila belum ada."""
    kolom = {str(ws.cell(1, c).value).strip(): c
             for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}
    for wajib in ["Kecamatan", "Kabupaten/Kota"] + KOLOM_CUACA:
        if wajib not in kolom:
            raise SystemExit(f"Kolom '{wajib}' tidak ada di baris pertama Excel.")
    for nama in KOLOM_TAMBAHAN:
        if nama not in kolom:
            kolom[nama] = ws.max_column + 1
            sel = ws.cell(1, kolom[nama], nama)
            sel.font = Font(bold=True, color="FFFFFF", name="Arial")
            sel.fill = PatternFill("solid", fgColor="1F4E78")
            sel.alignment = Alignment(horizontal="center", vertical="center")
    return kolom


def baca_tugas(ws, kolom):
    """Kumpulkan baris yang belum terisi (baris sudah terisi dilewati -> bisa dijalankan ulang)."""
    tugas, dilewati = [], 0
    for baris in range(2, ws.max_row + 1):
        kec = ws.cell(baris, kolom["Kecamatan"]).value
        if not kec or not str(kec).strip():
            continue
        if ws.cell(baris, kolom["Last Update"]).value not in (None, "", "-"):
            dilewati += 1
            continue
        kab = ws.cell(baris, kolom["Kabupaten/Kota"]).value or ""
        lat = ws.cell(baris, kolom["Latitude"]).value
        lon = ws.cell(baris, kolom["Longitude"]).value
        try:
            lat = float(lat) if lat not in (None, "") else None
            lon = float(lon) if lon not in (None, "") else None
        except ValueError:
            lat = lon = None
        tugas.append((baris, str(kab).strip(), str(kec).strip(), lat, lon))
    return tugas, dilewati


def tulis_hasil(ws, kolom, hasil):
    for baris, d in hasil.items():
        if "error" in d:  # tandai gagal dengan '-' dan simpan pesan galatnya
            for nama in KOLOM_CUACA:
                ws.cell(baris, kolom[nama], "-")
            ws.cell(baris, kolom["Catatan"], d["error"])
        else:
            for nama in KOLOM_CUACA:
                ws.cell(baris, kolom[nama], d[nama])
            ws.cell(baris, kolom["Catatan"], d.get("Catatan", ""))
        ws.cell(baris, kolom["Latitude"], d.get("Latitude"))
        ws.cell(baris, kolom["Longitude"], d.get("Longitude"))

    for nama, lebar in {"Last Update": 22, "Kondisi Cuaca": 22, "Kecepatan Angin (km/jam)": 22,
                        "Arah Angin": 20, "Catatan": 40}.items():
        ws.column_dimensions[ws.cell(1, kolom[nama]).column_letter].width = lebar


# ============================== PROGRAM UTAMA ==============================
def main():
    file_excel = sys.argv[1] if len(sys.argv) > 1 else FILE_EXCEL
    jumlah_thread = int(sys.argv[2]) if len(sys.argv) > 2 else JUMLAH_THREAD

    wb = load_workbook(file_excel)
    ws = wb.active
    kolom = peta_kolom(ws)
    tugas, dilewati = baca_tugas(ws, kolom)
    if not tugas:
        print(f"Tidak ada tugas baru ({dilewati} baris sudah terisi).")
        return
    print(f"Kecamatan diproses: {len(tugas)} (dilewati: {dilewati}) | Thread: {jumlah_thread}")

    # Producer: isi antrean dengan semua tugas
    antrean = queue.Queue()
    for t in tugas:
        antrean.put(t)
    hasil, lock = {}, threading.Lock()
    progres = {"selesai": 0, "total": len(tugas)}

    # Consumer: jalankan semua thread, tunggu sampai selesai (join)
    mulai = time.time()
    pekerja = [PekerjaCuaca(i + 1, antrean, hasil, lock, progres) for i in range(jumlah_thread)]
    for p in pekerja:
        p.start()
    for p in pekerja:
        p.join()
    durasi = time.time() - mulai

    tulis_hasil(ws, kolom, hasil)
    try:
        wb.save(file_excel)
    except PermissionError:  # file sedang dibuka di Excel
        file_excel = file_excel.replace(".xlsx", "_hasil.xlsx")
        wb.save(file_excel)
        print("File asli sedang dibuka di Excel, hasil disimpan ke file baru.")

    gagal = sum(1 for d in hasil.values() if "error" in d)
    pendekatan = sum(1 for d in hasil.values() if "pendekatan" in d.get("Catatan", ""))
    print(f"\nSelesai {durasi:.1f} detik | berhasil: {len(hasil) - gagal} | gagal: {gagal} | "
          f"koordinat pendekatan kab/kota: {pendekatan}")
    print(f"Hasil tersimpan di: {file_excel}")
    if gagal:
        print("Jalankan sekali lagi untuk mengulang baris yang gagal.")


if __name__ == "__main__":
    main()